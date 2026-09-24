"""Server-side text extraction for uploaded deviation documents."""

import base64
import json

from fastapi import HTTPException
from langchain_core.messages import HumanMessage

from ..config import settings
from .llm import get_llm, get_vision_llm

MAX_FILE_BYTES = 10 * 1024 * 1024  # 10 MB, matches the UI copy

SUPPORTED_EXTENSIONS = {
    ".pdf": "PDF",
    ".docx": "DOCX",
    ".txt": "TXT",
    ".md": "TXT",
    ".xls": "XLS",
    ".xlsx": "XLS",
    ".csv": "XLS",
    ".jpg": "Image",
    ".jpeg": "Image",
    ".png": "Image",
}


def _decode(data: bytes) -> str:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _from_pdf(data: bytes) -> str:
    from io import BytesIO

    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _from_docx(data: bytes) -> str:
    from io import BytesIO

    from docx import Document

    doc = Document(BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _from_spreadsheet(data: bytes, filename: str) -> str:
    from io import BytesIO

    if filename.lower().endswith(".csv"):
        return _decode(data)

    if filename.lower().endswith(".xls"):
        try:
            import xlrd
        except ImportError:
            raise HTTPException(
                422, "Legacy .xls files are not supported. Please save as .xlsx or paste the text."
            )
        book = xlrd.open_workbook(file_contents=data)
        lines = []
        for sheet in book.sheets():
            for row in range(sheet.nrows):
                lines.append(" | ".join(str(v) for v in sheet.row_values(row)))
        return "\n".join(lines)

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(v) for v in row if v is not None]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


async def _from_image(data: bytes, filename: str) -> str:
    """OCR a photo/scan of a deviation report using a Groq vision model."""
    if not settings.groq_api_key:
        raise HTTPException(
            422,
            "Image OCR needs a Groq API key configured on the server. Please paste the text instead.",
        )
    b64 = base64.b64encode(data).decode()
    mime = "image/png" if filename.lower().endswith(".png") else "image/jpeg"
    llm = get_vision_llm()
    try:
        message = HumanMessage(
            content=[
                {"type": "text", "text": (
                    "Transcribe every word of this document verbatim. "
                    "Output only the transcribed text, nothing else."
                )},
                {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}},
            ]
        )
        response = await llm.ainvoke([message])
        return response.content if isinstance(response.content, str) else json.dumps(response.content)
    except Exception as exc:  # pragma: no cover - depends on Groq availability
        raise HTTPException(422, f"Could not read the image: {exc}")


async def extract_text_from_upload(filename: str, data: bytes) -> str:
    if len(data) > MAX_FILE_BYTES:
        raise HTTPException(413, "File exceeds the 10 MB limit.")

    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            422,
            f"Unsupported file type '{ext or filename}'. "
            "Supported formats: PDF, DOCX, TXT, XLS, JPG, PNG.",
        )

    try:
        if ext == ".pdf":
            text = _from_pdf(data)
        elif ext == ".docx":
            text = _from_docx(data)
        elif ext in (".xls", ".xlsx", ".csv"):
            text = _from_spreadsheet(data, filename)
        elif ext in (".jpg", ".jpeg", ".png"):
            text = await _from_image(data, filename)
        else:
            text = _decode(data)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, f"Could not read '{filename}': {exc}")

    if not text or not text.strip():
        raise HTTPException(
            422,
            "No readable text found in the file. If it is a scanned image, "
            "try a clearer photo or paste the text instead.",
        )
    return text.strip()
